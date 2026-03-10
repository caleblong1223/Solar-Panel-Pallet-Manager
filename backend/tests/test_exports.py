from __future__ import annotations

from collections.abc import Generator
from contextlib import contextmanager
from datetime import datetime, timezone

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401
from app.api.v1.endpoints.auth import get_current_user
from app.db.base import Base
from app.db.session import get_db
from app.main import app
from app.models.pallet import Export, Pallet, PalletItem, SimImportBatch, SimPanel


class DummyRole:
    def __init__(self, name: str) -> None:
        self.name = name


class DummyUser:
    def __init__(self, user_id: int, roles: list[str]) -> None:
        self.id = user_id
        self.roles = [DummyRole(role) for role in roles]


@contextmanager
def _client_with_pallet(
    user: DummyUser,
    *,
    pallet_status: str = "completed",
    max_panels: int = 25,
) -> Generator[TestClient, None, None]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    testing_session_local = sessionmaker(autocommit=False, autoflush=False, bind=engine, class_=Session)
    Base.metadata.create_all(bind=engine)

    seed_db = testing_session_local()
    try:
        pallet = Pallet(
            pallet_number=9,
            status=pallet_status,
            max_panels=max_panels,
            created_by=user.id,
            completed_at=datetime.now(timezone.utc) if pallet_status == "completed" else None,
        )
        seed_db.add(pallet)
        seed_db.flush()
        seed_db.add_all(
            [
                PalletItem(pallet_id=pallet.id, serial="EXP001", slot_index=1, added_by=user.id),
                PalletItem(pallet_id=pallet.id, serial="EXP002", slot_index=2, added_by=user.id),
            ]
        )
        batch = SimImportBatch(source_filename="sim.csv", status="completed")
        seed_db.add(batch)
        seed_db.flush()
        seed_db.add_all(
            [
                SimPanel(
                    batch_id=batch.id,
                    serial="EXP001",
                    watts=450.12,
                    isc=11.234,
                    voc=49.321,
                    imp=10.456,
                    vmp=40.111,
                    ff=0.812,
                ),
                SimPanel(
                    batch_id=batch.id,
                    serial="EXP002",
                    watts=451.01,
                    isc=11.101,
                    voc=49.101,
                    imp=10.300,
                    vmp=40.000,
                    ff=0.800,
                ),
            ]
        )
        seed_db.commit()
    finally:
        seed_db.close()

    def override_get_db() -> Generator[Session, None, None]:
        db = testing_session_local()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = lambda: user
    try:
        with TestClient(app) as client:
            yield client
    finally:
        app.dependency_overrides.clear()


def test_create_export_success(monkeypatch) -> None:
    with _client_with_pallet(DummyUser(user_id=1, roles=["packout_operator"])) as client:
        from app.api.v1.endpoints import exports as exports_endpoint

        monkeypatch.setattr(
            exports_endpoint,
            "_render_pdf_from_workbook_bytes",
            lambda workbook_bytes: b"pdf-from-workbook",
        )
        monkeypatch.setattr(
            exports_endpoint,
            "upload_export_artifact",
            lambda export_id, pallet_id, filename, content: (
                f"exports/2026/03/{pallet_id}/{export_id}/{filename}",
                "checksum123",
            ),
        )
        response = client.post("/api/v1/exports", json={"pallet_id": 1, "template_type": "450WT"})
        assert response.status_code == 201
        payload = response.json()
        assert payload["pallet_id"] == 1
        assert payload["checksum_sha256"] == "checksum123"
        assert payload["object_key"].startswith("exports/")


def test_create_export_requires_completed_pallet(monkeypatch) -> None:
    with _client_with_pallet(DummyUser(user_id=2, roles=["admin"]), pallet_status="active") as client:
        from app.api.v1.endpoints import exports as exports_endpoint

        monkeypatch.setattr(
            exports_endpoint,
            "_render_pdf_from_workbook_bytes",
            lambda workbook_bytes: b"pdf-from-workbook",
        )
        monkeypatch.setattr(
            exports_endpoint,
            "upload_export_artifact",
            lambda export_id, pallet_id, filename, content: (
                f"exports/2026/03/{pallet_id}/{export_id}/{filename}",
                "checksum123",
            ),
        )
        response = client.post("/api/v1/exports", json={"pallet_id": 1, "template_type": "450WT"})
        assert response.status_code == 409


def test_list_exports_and_filters(monkeypatch) -> None:
    with _client_with_pallet(DummyUser(user_id=3, roles=["purchasing_manager"])) as client:
        from app.api.v1.endpoints import exports as exports_endpoint

        monkeypatch.setattr(
            exports_endpoint,
            "_render_pdf_from_workbook_bytes",
            lambda workbook_bytes: b"pdf-from-workbook",
        )
        monkeypatch.setattr(
            exports_endpoint,
            "upload_export_artifact",
            lambda export_id, pallet_id, filename, content: (
                f"exports/2026/03/{pallet_id}/{export_id}/{filename}",
                f"checksum{export_id}",
            ),
        )
        create_a = client.post("/api/v1/exports", json={"pallet_id": 1, "template_type": "450WT"})
        create_b = client.post("/api/v1/exports", json={"pallet_id": 1, "template_type": "550WT"})
        assert create_a.status_code == 201
        assert create_b.status_code == 201

        list_all = client.get("/api/v1/exports")
        assert list_all.status_code == 200
        assert list_all.json()["total"] == 2

        filtered = client.get("/api/v1/exports?template_type=450WT")
        assert filtered.status_code == 200
        assert filtered.json()["total"] == 1
        assert filtered.json()["exports"][0]["template_type"] == "450WT"


def test_export_download_url_endpoint(monkeypatch) -> None:
    with _client_with_pallet(DummyUser(user_id=5, roles=["admin"])) as client:
        from app.api.v1.endpoints import exports as exports_endpoint

        monkeypatch.setattr(
            exports_endpoint,
            "_render_pdf_from_workbook_bytes",
            lambda workbook_bytes: b"pdf-from-workbook",
        )
        monkeypatch.setattr(
            exports_endpoint,
            "upload_export_artifact",
            lambda export_id, pallet_id, filename, content: (
                f"exports/2026/03/{pallet_id}/{export_id}/{filename}",
                f"checksum{export_id}",
            ),
        )
        created = client.post("/api/v1/exports", json={"pallet_id": 1, "template_type": "450WT"})
        export_id = created.json()["id"]

        monkeypatch.setattr(
            exports_endpoint,
            "generate_export_download_url",
            lambda object_key, expires_in_seconds=900: f"https://minio.local/{object_key}?exp={expires_in_seconds}",
        )
        dl = client.get(f"/api/v1/exports/{export_id}/download-url?expires_in_seconds=600")
        assert dl.status_code == 200
        assert dl.json()["export_id"] == export_id
        assert "https://minio.local/" in dl.json()["download_url"]

        missing = client.get("/api/v1/exports/999/download-url")
        assert missing.status_code == 404


def test_replace_export_workbook_endpoint(monkeypatch) -> None:
    with _client_with_pallet(DummyUser(user_id=6, roles=["packout_operator"])) as client:
        from app.api.v1.endpoints import exports as exports_endpoint

        monkeypatch.setattr(
            exports_endpoint,
            "_render_pdf_from_workbook_bytes",
            lambda workbook_bytes: b"pdf-from-workbook",
        )
        monkeypatch.setattr(
            exports_endpoint,
            "upload_export_artifact",
            lambda export_id, pallet_id, filename, content: (
                f"exports/2026/03/{pallet_id}/{export_id}/{filename}",
                f"checksum{export_id}",
            ),
        )
        created = client.post("/api/v1/exports", json={"pallet_id": 1, "template_type": "450WT"})
        export_id = created.json()["id"]

        captured: dict[str, object] = {}

        def _upload_at_key(*, object_key: str, filename: str, content: bytes, content_type: str | None = None):
            captured[filename] = {
                "object_key": object_key,
                "content_type": content_type,
                "content_len": len(content),
            }
            return object_key, "updatedchecksum"

        monkeypatch.setattr(exports_endpoint, "upload_export_artifact_at_key", _upload_at_key)
        files = {
            "file": (
                "edited.xlsx",
                b"edited-xlsx-bytes",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        replace = client.post(f"/api/v1/exports/{export_id}/replace", files=files)
        assert replace.status_code == 200
        payload = replace.json()
        assert payload["id"] == export_id
        xlsx_name = payload["file_name"].replace(".pdf", ".xlsx")
        assert xlsx_name in captured
        assert payload["file_name"] in captured
        assert str(captured[xlsx_name]["object_key"]).endswith(xlsx_name)
        assert captured[xlsx_name]["content_type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert captured[xlsx_name]["content_len"] == len(b"edited-xlsx-bytes")
        assert captured[payload["file_name"]]["content_type"] == "application/pdf"
        assert captured[payload["file_name"]]["content_len"] == len(b"pdf-from-workbook")


def test_create_export_populates_sim_values_in_workbook(monkeypatch) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    with _client_with_pallet(DummyUser(user_id=7, roles=["packout_operator"]), max_panels=25) as client:
        from app.api.v1.endpoints import exports as exports_endpoint

        monkeypatch.setattr(
            exports_endpoint,
            "_render_pdf_from_workbook_bytes",
            lambda workbook_bytes: b"pdf-from-workbook",
        )
        captured: dict[str, bytes] = {}

        def _upload_capture(export_id: int, pallet_id: int, filename: str, content: bytes):
            if filename.lower().endswith(".xlsx"):
                captured["xlsx"] = content
            return f"exports/2026/03/{pallet_id}/{export_id}/{filename}", "checksum123"

        monkeypatch.setattr(exports_endpoint, "upload_export_artifact", _upload_capture)
        response = client.post("/api/v1/exports", json={"pallet_id": 1, "template_type": "450WT"})
        assert response.status_code == 201
        assert "xlsx" in captured

        workbook = load_workbook(BytesIO(captured["xlsx"]))
        try:
            sheet = workbook["PALLET SHEET"]
            # Row 5 corresponds to EXP001 in seeded pallet items.
            assert sheet["B5"].value == "EXP001"
            assert float(sheet["C5"].value) == 450.12
            assert float(sheet["D5"].value) == 11.23
            assert float(sheet["E5"].value) == 49.32
            assert float(sheet["F5"].value) == 10.46
            assert float(sheet["G5"].value) == 40.11
        finally:
            workbook.close()


def test_create_export_prefers_in_spec_sim_data_for_template(monkeypatch) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    with _client_with_pallet(DummyUser(user_id=8, roles=["packout_operator"]), max_panels=25) as client:
        from app.api.v1.endpoints import exports as exports_endpoint

        monkeypatch.setattr(
            exports_endpoint,
            "_render_pdf_from_workbook_bytes",
            lambda workbook_bytes: b"pdf-from-workbook",
        )
        captured: dict[str, bytes] = {}

        def _upload_capture(export_id: int, pallet_id: int, filename: str, content: bytes):
            if filename.lower().endswith(".xlsx"):
                captured["xlsx"] = content
            return f"exports/2026/03/{pallet_id}/{export_id}/{filename}", "checksum123"

        monkeypatch.setattr(exports_endpoint, "upload_export_artifact", _upload_capture)

        db = next(iter(app.dependency_overrides[get_db]()))
        try:
            batch = SimImportBatch(source_filename="sim-extra.csv", status="completed")
            db.add(batch)
            db.flush()
            db.add(
                SimPanel(
                    batch_id=batch.id,
                    serial="EXP001",
                    panel_type="WRONGTYPE",
                    watts=470.50,
                    isc=11.999,
                    voc=50.999,
                    imp=10.999,
                    vmp=42.999,
                    ff=0.900,
                )
            )
            db.commit()
        finally:
            db.close()

        response = client.post("/api/v1/exports", json={"pallet_id": 1, "template_type": "450WT"})
        assert response.status_code == 201

        workbook = load_workbook(BytesIO(captured["xlsx"]))
        try:
            sheet = workbook["PALLET SHEET"]
            assert sheet["B5"].value == "EXP001"
            assert float(sheet["C5"].value) == 450.12
        finally:
            workbook.close()
