from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import platform
from pathlib import Path
import shutil
import subprocess
import tempfile

from openpyxl import load_workbook

from app.models.pallet import Pallet


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _content_stream_for_lines(lines: list[str]) -> bytes:
    y = 760
    text_ops = ["BT", "/F1 11 Tf", "72 0 0 72 0 0 Tm"]
    for line in lines:
        safe = _pdf_escape(line)
        text_ops.append(f"1 0 0 1 40 {y} Tm ({safe}) Tj")
        y -= 16
    text_ops.append("ET")
    return "\n".join(text_ops).encode("latin-1", errors="ignore")


def _pdf_from_pages(pages: list[list[str]]) -> bytes:
    if not pages:
        pages = [["No data"]]

    catalog_obj = 1
    pages_obj = 2
    font_obj = 3
    first_page_obj = 4

    objects: list[bytes] = []
    page_object_ids: list[int] = []
    content_object_ids: list[int] = []

    for index, lines in enumerate(pages):
        page_obj_id = first_page_obj + (index * 2)
        content_obj_id = page_obj_id + 1
        page_object_ids.append(page_obj_id)
        content_object_ids.append(content_obj_id)

    kids = " ".join(f"{obj_id} 0 R" for obj_id in page_object_ids)
    objects.append(
        f"{catalog_obj} 0 obj << /Type /Catalog /Pages {pages_obj} 0 R >> endobj\n".encode("ascii")
    )
    objects.append(
        f"{pages_obj} 0 obj << /Type /Pages /Count {len(page_object_ids)} /Kids [{kids}] >> endobj\n".encode(
            "ascii"
        )
    )
    objects.append(
        f"{font_obj} 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n".encode("ascii")
    )
    for idx, lines in enumerate(pages):
        page_obj_id = page_object_ids[idx]
        content_obj_id = content_object_ids[idx]
        stream = _content_stream_for_lines(lines)
        objects.append(
            f"{page_obj_id} 0 obj << /Type /Page /Parent {pages_obj} 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_obj} 0 R >> >> /Contents {content_obj_id} 0 R >> endobj\n".encode(
                "ascii"
            )
        )
        objects.append(
            f"{content_obj_id} 0 obj << /Length {len(stream)} >> stream\n".encode("ascii")
            + stream
            + b"\nendstream endobj\n"
        )

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(pdf))
        pdf.extend(obj)
    xref_pos = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode(
            "ascii"
        )
    )
    return bytes(pdf)


def generate_export_pdf_bytes(pallet: Pallet, template_type: str) -> bytes:
    now = datetime.now(timezone.utc).isoformat()
    lines = [
        "Pallet Manager Export",
        f"Pallet ID: {pallet.id}",
        f"Pallet Number: {pallet.pallet_number}",
        f"Template: {template_type}",
        f"Status: {pallet.status}",
        f"Created: {now}",
    ]
    if pallet.items:
        lines.append("Serials:")
        lines.extend([f" - {item.serial}" for item in pallet.items])
    return _pdf_from_pages([lines])


def _workbook_lines(workbook_bytes: bytes, title: str | None = None) -> list[str]:
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    try:
        sheet = workbook["PALLET SHEET"] if "PALLET SHEET" in workbook.sheetnames else workbook[workbook.sheetnames[0]]

        def _val(cell: str) -> str:
            value = sheet[cell].value
            if value is None:
                return ""
            if isinstance(value, datetime):
                return value.strftime("%Y-%m-%d")
            return str(value)

        lines = [
            title or "Pallet Sheet Export",
            f"Panel Model: {_val('B1')}",
            f"Quantity: {_val('B2')}",
            f"Weight (lbs): {_val('D2')}",
            f"Pallet Ref: {_val('B3')}",
            f"Packout Date: {_val('G3')}",
            "",
            "Serial  Pm  Isc  Voc  Ipm  Vpm",
        ]

        for row in range(5, 36):
            serial = sheet.cell(row=row, column=2).value
            if serial is None or str(serial).strip() == "":
                continue
            pm = sheet.cell(row=row, column=3).value
            isc = sheet.cell(row=row, column=4).value
            voc = sheet.cell(row=row, column=5).value
            ipm = sheet.cell(row=row, column=6).value
            vpm = sheet.cell(row=row, column=7).value
            lines.append(
                f"{str(serial).strip()}  {pm if pm is not None else ''}  {isc if isc is not None else ''}  {voc if voc is not None else ''}  {ipm if ipm is not None else ''}  {vpm if vpm is not None else ''}"
            )
        return lines
    finally:
        workbook.close()


def generate_pdf_from_workbook_bytes(workbook_bytes: bytes) -> bytes:
    return _pdf_from_pages([_workbook_lines(workbook_bytes)])


def generate_merged_pdf_from_workbook_pages(pages: list[tuple[str, bytes]]) -> bytes:
    rendered = [_workbook_lines(workbook_bytes, title=title) for title, workbook_bytes in pages]
    return _pdf_from_pages(rendered)


def convert_workbook_bytes_to_pdf_bytes(workbook_bytes: bytes) -> bytes:
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)
        source = tmp_path / "export.xlsx"
        output_pdf = tmp_path / "export.pdf"
        source.write_bytes(workbook_bytes)

        # 1) Prefer LibreOffice on all platforms when available.
        soffice_path = shutil.which("soffice")
        if soffice_path is None:
            mac_soffice = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
            if mac_soffice.exists():
                soffice_path = str(mac_soffice)

        if soffice_path is not None:
            process = subprocess.run(
                [
                    soffice_path,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(tmp_path),
                    str(source),
                ],
                capture_output=True,
                text=True,
                check=False,
            )
            if process.returncode == 0 and output_pdf.exists():
                return output_pdf.read_bytes()
            stderr = (process.stderr or "").strip()
            stdout = (process.stdout or "").strip()
            message = stderr or stdout or f"LibreOffice exited with code {process.returncode}"
            # On Windows, allow Excel fallback if LibreOffice conversion fails.
            if platform.system() != "Windows":
                raise RuntimeError(f"XLSX to PDF conversion failed: {message}")

        # 2) Windows fallback: Excel COM automation if Excel is installed.
        if platform.system() == "Windows":
            ps_script = r"""
param([string]$xlsxPath,[string]$pdfPath)
$excel = $null
$workbook = $null
try {
  $excel = New-Object -ComObject Excel.Application
  $excel.Visible = $false
  $excel.DisplayAlerts = $false
  $workbook = $excel.Workbooks.Open($xlsxPath)
  # xlTypePDF = 0
  $workbook.ExportAsFixedFormat(0, $pdfPath)
}
finally {
  if ($workbook -ne $null) { $workbook.Close($false) | Out-Null }
  if ($excel -ne $null) { $excel.Quit() | Out-Null }
}
"""
            process = subprocess.run(
                [
                    "powershell",
                    "-NoProfile",
                    "-NonInteractive",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-Command",
                    ps_script,
                    str(source),
                    str(output_pdf),
                ],
                capture_output=True,
                text=True,
                check=False,
            )
            if process.returncode == 0 and output_pdf.exists():
                return output_pdf.read_bytes()
            stderr = (process.stderr or "").strip()
            stdout = (process.stdout or "").strip()
            message = stderr or stdout or f"Excel conversion exited with code {process.returncode}"
            raise RuntimeError(f"XLSX to PDF conversion failed via LibreOffice/Excel: {message}")

        raise RuntimeError("XLSX to PDF conversion failed: LibreOffice (soffice) not found")
