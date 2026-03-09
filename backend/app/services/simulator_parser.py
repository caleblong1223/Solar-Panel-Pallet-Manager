from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime

from openpyxl import load_workbook


@dataclass
class ParsedSimRow:
    row_number: int
    serial: str
    panel_type: str | None
    result: str | None
    test_timestamp: datetime | None
    watts: float | None
    voc: float | None
    isc: float | None
    vmp: float | None
    imp: float | None
    ff: float | None


@dataclass
class RejectedSimRow:
    row_number: int
    reason: str
    raw_serial: str | None


@dataclass
class SimulatorParseResult:
    rows_total: int
    accepted_rows: list[ParsedSimRow]
    rejected_rows: list[RejectedSimRow]


def _normalize_key(value: str) -> str:
    # Normalize varied Sun Simulator headers like "Pmax(W)", "Date/Time", "Voc (V)".
    return re.sub(r"[^a-z0-9]+", "", value.strip().lower())


def _normalize_serial(value: object) -> str:
    return str(value).strip().upper()


def _find_column(headers: list[str], aliases: set[str]) -> int | None:
    normalized = [_normalize_key(h) for h in headers]
    for idx, key in enumerate(normalized):
        if key in aliases:
            return idx
    return None


def _find_column_contains(headers: list[str], fragments: tuple[str, ...]) -> int | None:
    normalized = [_normalize_key(h) for h in headers]
    for idx, key in enumerate(normalized):
        if any(fragment in key for fragment in fragments):
            return idx
    return None


def _parse_datetime(value: object) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    as_text = str(value).strip()
    if not as_text:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y"):
        try:
            return datetime.strptime(as_text, fmt)
        except ValueError:
            continue
    return None


def _parse_datetime_from_date_time(date_value: object, time_value: object) -> datetime | None:
    date_text = str(date_value).strip() if date_value is not None else ""
    time_text = str(time_value).strip() if time_value is not None else ""
    if not date_text and not time_text:
        return None

    combined = f"{date_text} {time_text}".strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%Y-%m-%d %I:%M:%S %p",
        "%m/%d/%Y %I:%M:%S %p",
        "%Y-%m-%d %I:%M %p",
        "%m/%d/%Y %I:%M %p",
    ):
        try:
            return datetime.strptime(combined, fmt)
        except ValueError:
            continue

    return _parse_datetime(combined)


def _parse_float(value: object) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    # Accept values like "49.321V", "0.812%", "1,234.56", etc.
    cleaned = text.replace(",", "")
    match = re.search(r"[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?", cleaned)
    if match:
        try:
            return float(match.group(0))
        except ValueError:
            pass
    try:
        return float(text)
    except ValueError:
        return None


def _build_result(headers: list[str], rows: list[list[object]]) -> SimulatorParseResult:
    serial_col = _find_column(
        headers,
        aliases={"serial", "serialno", "serialnumber", "barcodeserial", "barcode", "barcodeno", "modulesn"},
    )
    panel_type_col = _find_column(headers, aliases={"paneltype", "type", "moduletype"})
    result_col = _find_column(headers, aliases={"result", "status", "passfail"})
    ts_col = _find_column(headers, aliases={"testtimestamp", "timestamp", "testtime", "datetime", "dateandtime"})
    if ts_col is None:
        ts_col = _find_column_contains(headers, ("timestamp", "datetime", "date", "time"))
    date_col = _find_column(headers, aliases={"date", "testdate"})
    time_col = _find_column(headers, aliases={"time", "testtime", "testhour"})

    watts_col = _find_column(headers, aliases={"watts", "watt", "pm", "pmax", "pmaxw"})
    if watts_col is None:
        watts_col = _find_column_contains(headers, ("pmax", "watts", "watt", "power", "pm"))

    voc_col = _find_column(headers, aliases={"voc", "vocv"})
    if voc_col is None:
        voc_col = _find_column_contains(headers, ("voc",))

    isc_col = _find_column(headers, aliases={"isc", "isca"})
    if isc_col is None:
        isc_col = _find_column_contains(headers, ("isc",))

    vmp_col = _find_column(headers, aliases={"vmp", "vmpp", "vpm", "vmpv", "vpmv"})
    if vmp_col is None:
        vmp_col = _find_column_contains(headers, ("vmp", "vmpp", "vpm"))

    imp_col = _find_column(headers, aliases={"imp", "impp", "ipm", "impa"})
    if imp_col is None:
        imp_col = _find_column_contains(headers, ("imp", "impp"))

    ff_col = _find_column(headers, aliases={"ff", "fillfactor"})
    if ff_col is None:
        ff_col = _find_column_contains(headers, ("fillfactor", "ff"))

    accepted: list[ParsedSimRow] = []
    rejected: list[RejectedSimRow] = []
    for idx, row in enumerate(rows, start=2):
        if serial_col is None or serial_col >= len(row):
            rejected.append(RejectedSimRow(row_number=idx, reason="Missing serial column", raw_serial=None))
            continue
        serial = _normalize_serial(row[serial_col])
        if not serial:
            rejected.append(RejectedSimRow(row_number=idx, reason="Empty serial", raw_serial=None))
            continue
        test_ts = _parse_datetime(row[ts_col]) if ts_col is not None and ts_col < len(row) else None
        should_prefer_date_time = ts_col is None or ts_col == date_col or ts_col == time_col
        if (test_ts is None or should_prefer_date_time) and (date_col is not None or time_col is not None):
            date_value = row[date_col] if date_col is not None and date_col < len(row) else None
            time_value = row[time_col] if time_col is not None and time_col < len(row) else None
            combined_ts = _parse_datetime_from_date_time(date_value, time_value)
            if combined_ts is not None:
                test_ts = combined_ts
        panel_type = (
            str(row[panel_type_col]).strip() if panel_type_col is not None and panel_type_col < len(row) and row[panel_type_col] is not None else None
        )
        result = (
            str(row[result_col]).strip().upper() if result_col is not None and result_col < len(row) and row[result_col] is not None else None
        )
        accepted.append(
            ParsedSimRow(
                row_number=idx,
                serial=serial,
                panel_type=panel_type,
                result=result,
                test_timestamp=test_ts,
                watts=_parse_float(row[watts_col]) if watts_col is not None and watts_col < len(row) else None,
                voc=_parse_float(row[voc_col]) if voc_col is not None and voc_col < len(row) else None,
                isc=_parse_float(row[isc_col]) if isc_col is not None and isc_col < len(row) else None,
                vmp=_parse_float(row[vmp_col]) if vmp_col is not None and vmp_col < len(row) else None,
                imp=_parse_float(row[imp_col]) if imp_col is not None and imp_col < len(row) else None,
                ff=_parse_float(row[ff_col]) if ff_col is not None and ff_col < len(row) else None,
            )
        )

    return SimulatorParseResult(rows_total=len(rows), accepted_rows=accepted, rejected_rows=rejected)


def parse_simulator_file(filename: str, content: bytes) -> SimulatorParseResult:
    lowered = filename.lower()
    if lowered.endswith(".csv"):
        text_stream = io.StringIO(content.decode("utf-8-sig", errors="ignore"))
        reader = csv.reader(text_stream)
        all_rows = list(reader)
    elif lowered.endswith(".xlsx") or lowered.endswith(".xlsm"):
        workbook = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        all_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    else:
        raise ValueError("Unsupported file type. Use .csv, .xlsx, or .xlsm")

    if not all_rows:
        return SimulatorParseResult(rows_total=0, accepted_rows=[], rejected_rows=[])

    headers = [str(cell or "") for cell in all_rows[0]]
    data_rows = [list(row) for row in all_rows[1:]]
    return _build_result(headers, data_rows)
