from __future__ import annotations

from datetime import datetime
from io import BytesIO
import os
from pathlib import Path
import re

from openpyxl import load_workbook

from app.models.pallet import Pallet


class ExportWorkbookError(RuntimeError):
    pass


def _repo_root() -> Path:
    # backend/app/services -> repo root
    return Path(__file__).resolve().parents[3]


def _template_path_for_capacity(max_panels: int) -> Path:
    if max_panels in (25, 26):
        template_name = "26.xlsx"
    elif max_panels == 30:
        template_name = "30.xlsx"
    elif max_panels == 35:
        template_name = "35.xlsx"
    else:
        raise ExportWorkbookError(
            f"Unsupported panel capacity {max_panels}. Allowed capacities: 25, 26, 30, 35."
        )

    root = _repo_root()
    candidates = [
        root / "data" / "EXCEL" / template_name,
        root / "EXCEL" / template_name,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise ExportWorkbookError(
        f"Template workbook {template_name} not found. Checked: "
        + ", ".join(str(path) for path in candidates)
    )


def _is_testing() -> bool:
    """Detect pytest so we can skip hard failures during test runs."""
    return bool(os.getenv("PYTEST_CURRENT_TEST"))


def verify_core_export_templates() -> None:
    """Ensure the critical Excel templates exist and are structurally valid.

    This validates the 26/30/35 panel templates on startup so we fail fast if
    the workbooks are missing or corrupted. In test runs we skip this check.
    """
    if _is_testing():
        return

    # Core capacities we expect exact templates for.
    for capacity in (26, 30, 35):
        path = _template_path_for_capacity(capacity)
        try:
            workbook = load_workbook(path)
        except Exception as exc:  # pragma: no cover - defensive
            raise ExportWorkbookError(f"Failed to load template workbook {path}: {exc}") from exc

        try:
            if "PALLET SHEET" not in workbook.sheetnames:
                raise ExportWorkbookError(
                    f"Template {path.name} missing required 'PALLET SHEET' worksheet."
                )
        finally:
            workbook.close()


def inspect_core_export_templates() -> dict[str, dict[str, object]]:
    """Return a detailed, non-raising view of template status for debugging."""
    status: dict[str, dict[str, object]] = {}
    for capacity in (26, 30, 35):
        key = f"{capacity}"
        info: dict[str, object] = {
            "capacity": capacity,
            "path": None,
            "exists": False,
            "valid": False,
            "error": None,
        }
        try:
            path = _template_path_for_capacity(capacity)
            info["path"] = str(path)
            exists = path.exists()
            info["exists"] = exists
            if not exists:
                info["error"] = "Template file not found"
            else:
                try:
                    workbook = load_workbook(path)
                except Exception as exc:  # pragma: no cover - defensive
                    info["error"] = f"Failed to load workbook: {exc}"
                else:
                    try:
                        if "PALLET SHEET" not in workbook.sheetnames:
                            info["error"] = "Missing 'PALLET SHEET' worksheet"
                        else:
                            info["valid"] = True
                    finally:
                        workbook.close()
        except ExportWorkbookError as exc:
            info["error"] = str(exc)
        status[key] = info
    return status


def _build_b3_value(panel_type: str, pallet_number: int, export_dt: datetime) -> str:
    # Match 1.1 format: {PanelType}{MDYYYY}-{PalletNumber}
    # Example: 200WT162025-1 for Jan 6, 2025 pallet 1
    date_mdyyyy = f"{export_dt.month}{export_dt.day}{export_dt.year}"
    return f"{panel_type}{date_mdyyyy}-{pallet_number}"


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def _resolve_sim_columns(sheet) -> dict[str, int]:
    header_aliases: dict[str, tuple[str, ...]] = {
        "pm": ("pm", "pmw", "pmax", "pmaxw", "watts", "watt"),
        "isc": ("isc", "isca"),
        "voc": ("voc", "vocv"),
        "ipm": ("ipm", "impp", "imp", "impa", "ipma"),
        "vpm": ("vpm", "vpmv", "vmp", "vmpv", "vmpp", "vpma"),
    }
    resolved: dict[str, int] = {}
    for col in range(1, 32):
        key = _normalize_header(sheet.cell(row=4, column=col).value)
        if not key:
            continue
        for field, aliases in header_aliases.items():
            if field in resolved:
                continue
            if key in aliases:
                resolved[field] = col
                break
    return resolved


def _resolve_ff_column(sheet) -> int | None:
    ff_aliases = {"ff", "ffpercent", "fillfactor"}
    for col in range(1, 32):
        key = _normalize_header(sheet.cell(row=4, column=col).value)
        if key in ff_aliases:
            return col
    return None


def _round_electrical(value: float | None) -> float | None:
    if value is None:
        return None
    return round(float(value), 2)


def generate_export_workbook_bytes(
    pallet: Pallet,
    panel_type: str,
    export_dt: datetime | None = None,
    sim_values_by_serial: dict[str, dict[str, float | None]] | None = None,
) -> bytes:
    when = export_dt or datetime.now()
    excel_when = when.replace(tzinfo=None) if when.tzinfo is not None else when
    template_path = _template_path_for_capacity(pallet.max_panels)
    workbook = load_workbook(template_path)
    try:
        if "PALLET SHEET" not in workbook.sheetnames:
            raise ExportWorkbookError(f"Template {template_path.name} missing 'PALLET SHEET' worksheet.")
        sheet = workbook["PALLET SHEET"]

        # 1.1 compatibility cells.
        sheet["B1"] = panel_type
        sheet["B2"] = len(pallet.items)
        sheet["D2"] = len(pallet.items) * 40
        sheet["G3"] = excel_when
        sheet["B3"] = _build_b3_value(panel_type, pallet.pallet_number, when)

        # Default customer block in A3, matching 1.1 layout.
        # Single station customer:
        #   Name:    Josh Atwood
        #   Business: Future Solutions Inc
        #   Street:  2616 Glenview Dr
        #   City:    Elkhart, IN 46514
        customer_block = "\n".join(
            [
                "Josh Atwood",
                "Future Solutions Inc",
                "2616 Glenview Dr",
                "Elkhart, IN 46514",
            ]
        )
        sheet["A3"] = customer_block

        sim_values_by_serial = sim_values_by_serial or {}
        sim_columns = _resolve_sim_columns(sheet)
        ff_col = _resolve_ff_column(sheet)

        # Clear serial slots then write current pallet serials into B5..B30
        for row in range(5, 31):
            sheet.cell(row=row, column=2).value = None
            for col in sim_columns.values():
                sheet.cell(row=row, column=col).value = None
            if ff_col is not None:
                sheet.cell(row=row, column=ff_col).value = None
        if ff_col is not None:
            sheet.cell(row=4, column=ff_col).value = None
        for idx, item in enumerate(sorted(pallet.items, key=lambda value: value.slot_index), start=5):
            if idx > 30:
                break
            sheet.cell(row=idx, column=2).value = item.serial
            serial_key = (item.serial or "").strip().upper()
            values = sim_values_by_serial.get(serial_key)
            if not values:
                continue
            if "pm" in sim_columns:
                sheet.cell(row=idx, column=sim_columns["pm"]).value = _round_electrical(values.get("pm"))
            if "isc" in sim_columns:
                sheet.cell(row=idx, column=sim_columns["isc"]).value = _round_electrical(values.get("isc"))
            if "voc" in sim_columns:
                sheet.cell(row=idx, column=sim_columns["voc"]).value = _round_electrical(values.get("voc"))
            if "ipm" in sim_columns:
                sheet.cell(row=idx, column=sim_columns["ipm"]).value = _round_electrical(values.get("ipm"))
            if "vpm" in sim_columns:
                sheet.cell(row=idx, column=sim_columns["vpm"]).value = _round_electrical(values.get("vpm"))

        out = BytesIO()
        workbook.save(out)
        return out.getvalue()
    finally:
        workbook.close()
