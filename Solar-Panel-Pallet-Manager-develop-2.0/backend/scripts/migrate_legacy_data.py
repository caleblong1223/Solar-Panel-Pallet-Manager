#!/usr/bin/env python3
"""
Legacy data migration for Pallet Manager 2.0.

Migrates from current 1.1 file-based storage into Postgres tables:
- customers
- pallets
- pallet_items
- exports

Usage:
  python backend/scripts/migrate_legacy_data.py --dry-run
  python backend/scripts/migrate_legacy_data.py --apply

Environment:
  DATABASE_URL=postgresql://user:pass@host:5432/dbname
"""

from __future__ import annotations

import argparse
import json
import os
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CUSTOMERS_XLSX = ROOT / "data" / "CUSTOMERS" / "customers.xlsx"
DEFAULT_PALLET_HISTORY = ROOT / "data" / "PALLETS" / "pallet_history.json"


@dataclass
class Stats:
    customers_seen: int = 0
    customers_inserted: int = 0
    pallets_seen: int = 0
    pallets_inserted: int = 0
    items_inserted: int = 0
    exports_inserted: int = 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Migrate legacy Pallet Manager data to Postgres")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--dry-run", action="store_true", help="Analyze and print planned changes")
    mode.add_argument("--apply", action="store_true", help="Apply migration changes")
    parser.add_argument("--customers", type=Path, default=DEFAULT_CUSTOMERS_XLSX)
    parser.add_argument("--history", type=Path, default=DEFAULT_PALLET_HISTORY)
    parser.add_argument("--default-user-id", type=int, default=1)
    return parser.parse_args()


def parse_customer_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists() or load_workbook is None:
        return []

    wb = load_workbook(path, data_only=True)
    ws = wb.active

    headers = [str(cell.value).strip().lower() if cell.value else "" for cell in ws[1]]

    def idx(*names: str) -> int | None:
        for name in names:
            if name in headers:
                return headers.index(name)
        return None

    i_name = idx("name", "contact", "contact_name")
    i_business = idx("business", "business_name", "company")
    i_email = idx("email", "email address")
    i_phone = idx("phone", "phone number", "telephone")

    out: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[i_name]).strip() if i_name is not None and i_name < len(row) and row[i_name] else ""
        business = (
            str(row[i_business]).strip()
            if i_business is not None and i_business < len(row) and row[i_business]
            else ""
        )
        email = str(row[i_email]).strip() if i_email is not None and i_email < len(row) and row[i_email] else ""
        phone = str(row[i_phone]).strip() if i_phone is not None and i_phone < len(row) and row[i_phone] else ""

        if not name and not business:
            continue

        display_name = f"{name} | {business}".strip(" |") if (name and business) else (name or business)
        out.append(
            {
                "display_name": display_name,
                "contact_name": name or None,
                "business_name": business or None,
                "email": email or None,
                "phone": phone or None,
            }
        )

    return out


def parse_history(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8") as f:
        data = json.load(f)
    return data.get("pallets", [])


def parse_legacy_dt(raw: str | None) -> datetime | None:
    if not raw:
        return None

    candidates = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d",
    ]
    for fmt in candidates:
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def upsert_customer(cur: Any, row: dict[str, Any], apply: bool) -> int | None:
    if not apply:
        return None

    cur.execute(
        """
        INSERT INTO customers (display_name, contact_name, business_name, email, phone)
        VALUES (%s, %s, %s, %s, %s)
        ON CONFLICT (display_name)
        DO UPDATE SET
          contact_name = EXCLUDED.contact_name,
          business_name = EXCLUDED.business_name,
          email = EXCLUDED.email,
          phone = EXCLUDED.phone
        RETURNING id
        """,
        (
            row["display_name"],
            row["contact_name"],
            row["business_name"],
            row["email"],
            row["phone"],
        ),
    )
    return cur.fetchone()[0]


def get_customer_id(cur: Any, display_name: str | None) -> int | None:
    if not display_name:
        return None
    cur.execute("SELECT id FROM customers WHERE display_name = %s", (display_name,))
    row = cur.fetchone()
    return row[0] if row else None


def resolve_actor_user_id(cur: Any, requested_user_id: int | None) -> int | None:
    if requested_user_id is None:
        return None
    cur.execute("SELECT id FROM users WHERE id = %s", (requested_user_id,))
    row = cur.fetchone()
    return row[0] if row else None


def insert_pallet(cur: Any, pallet: dict[str, Any], user_id: int, apply: bool) -> tuple[int | None, bool]:
    completed_at = parse_legacy_dt(pallet.get("completed_at"))
    status = "completed" if completed_at else "active"
    customer_display = None

    customer_obj = pallet.get("customer")
    if isinstance(customer_obj, dict):
        customer_display = customer_obj.get("display_name")
    elif isinstance(customer_obj, str):
        customer_display = customer_obj.strip() or None

    customer_id = get_customer_id(cur, customer_display) if apply else None

    if not apply:
        return None, False

    cur.execute(
        """
        SELECT id FROM pallets
        WHERE pallet_number = %s
          AND status = %s
          AND COALESCE(template_type, '') = COALESCE(%s, '')
          AND COALESCE(customer_id, -1) = COALESCE(%s, -1)
          AND (
            (completed_at IS NULL AND %s IS NULL)
            OR completed_at = %s
          )
        ORDER BY id DESC
        LIMIT 1
        """,
        (
            int(pallet.get("pallet_number") or 0),
            status,
            pallet.get("panel_type"),
            customer_id,
            completed_at,
            completed_at,
        ),
    )
    existing = cur.fetchone()
    if existing:
        return existing[0], False

    cur.execute(
        """
        INSERT INTO pallets (
          pallet_number, status, template_type, max_panels,
          customer_id, created_by, completed_by, created_at, completed_at
        )
        VALUES (%s, %s, %s, %s, %s, %s, %s, CURRENT_TIMESTAMP, %s)
        RETURNING id
        """,
        (
            int(pallet.get("pallet_number") or 0),
            status,
            pallet.get("panel_type"),
            int(pallet.get("max_panels") or 25),
            customer_id,
            user_id,
            user_id if completed_at else None,
            completed_at,
        ),
    )
    return cur.fetchone()[0], True


def insert_pallet_items(cur: Any, pallet_id: int, serials: list[str], user_id: int, apply: bool) -> int:
    count = 0
    if not apply:
        return len(serials)
    for idx, serial in enumerate(serials):
        if not serial:
            continue
        cur.execute(
            """
            INSERT INTO pallet_items (pallet_id, serial, slot_index, added_by)
            VALUES (%s, %s, %s, %s)
            ON CONFLICT DO NOTHING
            """,
            (pallet_id, serial.strip().upper(), idx + 1, user_id),
        )
        if cur.rowcount > 0:
            count += 1
    return count


def insert_export(cur: Any, pallet_id: int, pallet: dict[str, Any], user_id: int, apply: bool) -> int:
    exported_file = pallet.get("exported_file")
    if not exported_file:
        return 0

    object_key = str(exported_file).replace("\\", "/")
    file_name = Path(object_key).name
    template_type = pallet.get("panel_type") or "legacy"

    if not apply:
        return 1

    cur.execute("SELECT id FROM exports WHERE pallet_id = %s AND object_key = %s", (pallet_id, object_key))
    if cur.fetchone():
        return 0

    cur.execute(
        """
        INSERT INTO exports (pallet_id, template_type, object_key, file_name, created_by)
        VALUES (%s, %s, %s, %s, %s)
        """,
        (pallet_id, template_type, object_key, file_name, user_id),
    )
    return 1


def main() -> None:
    args = parse_args()
    apply = args.apply

    customers = parse_customer_rows(args.customers)
    pallets = parse_history(args.history)

    print(f"Loaded {len(customers)} customer rows from {args.customers}")
    print(f"Loaded {len(pallets)} pallets from {args.history}")

    stats = Stats()

    if not apply:
        stats.customers_seen = len(customers)
        stats.customers_inserted = len(customers)
        stats.pallets_seen = len(pallets)
        stats.items_inserted = sum(len(p.get("serial_numbers") or []) for p in pallets)
        stats.exports_inserted = sum(1 for p in pallets if p.get("exported_file"))
    else:
        import psycopg

        db_url = os.getenv("DATABASE_URL")
        if not db_url:
            raise SystemExit("DATABASE_URL environment variable is required")

        with psycopg.connect(db_url) as conn:
            with conn.cursor() as cur:
                actor_user_id = resolve_actor_user_id(cur, args.default_user_id)
                if actor_user_id is None:
                    print("Default user id not found; migration will use NULL actor ids for created_by/completed_by/added_by.")

                for c in customers:
                    stats.customers_seen += 1
                    upsert_customer(cur, c, apply=apply)
                    stats.customers_inserted += 1

                for p in pallets:
                    stats.pallets_seen += 1
                    pallet_id, inserted = insert_pallet(cur, p, actor_user_id, apply=apply)
                    if inserted and pallet_id is not None:
                        stats.pallets_inserted += 1

                    serials = p.get("serial_numbers") or []
                    stats.items_inserted += insert_pallet_items(
                        cur,
                        pallet_id if pallet_id is not None else -1,
                        serials,
                        actor_user_id,
                        apply=apply,
                    )

                    stats.exports_inserted += insert_export(
                        cur,
                        pallet_id if pallet_id is not None else -1,
                        p,
                        actor_user_id,
                        apply=apply,
                    )

            conn.commit()

    mode = "APPLY" if apply else "DRY-RUN"
    print(f"\nMigration mode: {mode}")
    print("Summary:")
    print(f"- Customers seen: {stats.customers_seen}")
    print(f"- Customers upserted (planned/applied): {stats.customers_inserted}")
    print(f"- Pallets seen: {stats.pallets_seen}")
    print(f"- Pallets inserted (planned/applied): {stats.pallets_inserted if apply else stats.pallets_seen}")
    print(f"- Pallet items inserted (planned/applied): {stats.items_inserted}")
    print(f"- Export metadata inserted (planned/applied): {stats.exports_inserted}")


if __name__ == "__main__":
    main()
