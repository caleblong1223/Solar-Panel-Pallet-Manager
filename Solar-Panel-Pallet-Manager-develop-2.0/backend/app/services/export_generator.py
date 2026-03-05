from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.pallet import Customer, Pallet, SimPanel


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def generate_export_pdf_bytes(pallet: Pallet, template_type: str) -> bytes:
    now = datetime.now(timezone.utc).isoformat()
    lines = [
        "Pallet Manager Export",
        f"Pallet ID: {pallet.id}",
        f"Pallet Number: {pallet.pallet_number}",
        f"Template: {template_type}",
        f"Status: {pallet.status}",
        f"Created: {now}",
    ]
    if pallet.items:
        lines.append("Serials:")
        lines.extend([f" - {item.serial}" for item in pallet.items])

    y = 760
    text_ops = ["BT", "/F1 11 Tf", "72 0 0 72 0 0 Tm"]
    for line in lines:
        safe = _pdf_escape(line)
        text_ops.append(f"1 0 0 1 40 {y} Tm ({safe}) Tj")
        y -= 16
    text_ops.append("ET")
    stream = "\n".join(text_ops).encode("latin-1", errors="ignore")

    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        b"2 0 obj << /Type /Pages /Count 1 /Kids [3 0 R] >> endobj\n",
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj\n",
        b"4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
        f"5 0 obj << /Length {len(stream)} >> stream\n".encode("ascii")
        + stream
        + b"\nendstream endobj\n",
    ]

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(pdf))
        pdf.extend(obj)
    xref_pos = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode(
            "ascii"
        )
    )
    return bytes(pdf)


def generate_multi_pallets_pdf_bytes(pallets: Sequence[Pallet]) -> bytes:
  """Generate a minimal combined PDF for multiple pallets, optimized for quick printing.

  This reuses the same lightweight PDF approach as generate_export_pdf_bytes but
  concatenates human-readable sections for each pallet into a single document.
  """
  lines: list[str] = ["Pallet Manager Export (combined)", ""]
  now = datetime.now(timezone.utc).isoformat()
  for index, pallet in enumerate(pallets, start=1):
      lines.append(f"Pallet #{index}")
      lines.append(f"Pallet ID: {pallet.id}")
      lines.append(f"Pallet Number: {pallet.pallet_number}")
      lines.append(f"Status: {pallet.status}")
      lines.append(f"Created: {pallet.created_at.isoformat()}")
      if pallet.completed_at is not None:
          lines.append(f"Completed: {pallet.completed_at.isoformat()}")
      lines.append(f"Export Generated: {now}")
      if pallet.items:
          lines.append("Serials:")
          lines.extend([f" - {item.serial}" for item in sorted(pallet.items, key=lambda i: i.slot_index)])
      lines.append("")  # blank line between pallets

  y = 760
  text_ops = ["BT", "/F1 11 Tf", "72 0 0 72 0 0 Tm"]
  for line in lines:
      safe = _pdf_escape(line)
      text_ops.append(f"1 0 0 1 40 {y} Tm ({safe}) Tj")
      y -= 16
  text_ops.append("ET")
  stream = "\n".join(text_ops).encode("latin-1", errors="ignore")

  objects = [
      b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
      b"2 0 obj << /Type /Pages /Count 1 /Kids [3 0 R] >> endobj\n",
      b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj\n",
      b"4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
      f"5 0 obj << /Length {len(stream)} >> stream\n".encode("ascii")
      + stream
      + b"\nendstream endobj\n",
  ]

  pdf = bytearray(b"%PDF-1.4\n")
  offsets = [0]
  for obj in objects:
      offsets.append(len(pdf))
      pdf.extend(obj)
  xref_pos = len(pdf)
  pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
  pdf.extend(b"0000000000 65535 f \n")
  for offset in offsets[1:]:
      pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
  pdf.extend(
      f"trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode(
          "ascii"
      )
  )
  return bytes(pdf)


def _find_pallet_workbook(excel_dir: Path, max_panels: int | None = None) -> Path | None:
    """Find the pallet template workbook using 1.1-style rules, with 26/30/35 templates."""
    # Panel-count-specific templates: 26.xlsx, 30.xlsx, 35.xlsx
    if max_panels is not None:
        if max_panels in (25, 26):
            candidate = excel_dir / "26.xlsx"
            if candidate.exists():
                return candidate
        elif max_panels == 30:
            candidate = excel_dir / "30.xlsx"
            if candidate.exists():
                return candidate
        elif max_panels == 35:
            candidate = excel_dir / "35.xlsx"
            if candidate.exists():
                return candidate

    current = excel_dir / "CURRENT.xlsx"
    if current.exists():
        return current

    if not excel_dir.exists():
        return None

    build_files: list[Path] = []
    for file in excel_dir.glob("*.xlsx"):
        if file.name.startswith("~$"):
            continue
        name = file.name
        upper = name.upper()
        if "BUILD" in upper and any(ch.isdigit() for ch in name) and "Q" in upper:
            build_files.append(file)
    if not build_files:
        return None
    return max(build_files, key=lambda p: p.stat().st_mtime)


def _panel_pm_ranges() -> dict[str, tuple[float, float]]:
    return {
        "200WT": (195, 206),
        "220WT": (214, 227),
        "220M6": (214, 227),
        "330WT": (320, 340),
        "450WT": (439, 463.5),
        "450BT": (439, 463.5),
    }


def _generate_theoretical_electrical_values(panel_type: str | None) -> dict[str, float]:
    ranges = _panel_pm_ranges()
    pm_min, pm_max = 350.0, 450.0
    if panel_type and panel_type in ranges:
        pm_min, pm_max = ranges[panel_type]
    # Use a deterministic seed per panel_type so values are stable.
    import random as _random

    seed_source = panel_type or f"{pm_min}-{pm_max}"
    rnd = _random.Random(seed_source)
    pm = rnd.uniform(pm_min, pm_max)
    voc = rnd.uniform(38.0, 50.0)
    vpm = voc * rnd.uniform(0.75, 0.85)
    ipm = pm / vpm if vpm > 0 else rnd.uniform(8.0, 12.0)
    isc = ipm / rnd.uniform(0.90, 0.98)
    return {
        "Pm": round(pm, 2),
        "Isc": round(isc, 3),
        "Voc": round(voc, 3),
        "Ipm": round(ipm, 3),
        "Vpm": round(vpm, 3),
    }


def _validate_pm_range(pm_value: float | None, panel_type: str | None) -> bool:
    if pm_value is None or not panel_type:
        return True
    ranges = _panel_pm_ranges()
    if panel_type not in ranges:
        return True
    lo, hi = ranges[panel_type]
    return lo <= pm_value <= hi


def _find_column_by_header(sheet, header_variations: Iterable[str]) -> str | None:
    header_variations_lower = [v.lower() for v in header_variations]
    max_col = min(getattr(sheet, "max_column", 26) or 26, 26)
    max_row = getattr(sheet, "max_row", 100) or 100
    for row in range(1, min(6, max_row + 1)):
        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            cell = sheet[f"{col_letter}{row}"]
            if cell.value:
                cell_value = str(cell.value).strip().lower()
                for variation in header_variations_lower:
                    if variation in cell_value or cell_value in variation:
                        return col_letter
    return None


def _update_pallet_sheet_from_db(sheet, pallet: Pallet, template_type: str, customer: Customer | None, serial_data_cache: dict[str, dict], export_dt: datetime) -> None:
    # Customer in A3 exactly as 1.1: name | business, then address and city/state/zip
    if customer:
        name_line = customer.display_name
        if customer.business_name:
            name_line = f"{customer.display_name} | {customer.business_name}"
        address_line = customer.address or ""
        city_state_zip = ""
        city_parts = []
        if customer.city:
            city_parts.append(customer.city)
        if customer.state:
            city_parts.append(customer.state)
        if customer.zip_code:
            city_parts.append(customer.zip_code)
        if city_parts:
            if len(city_parts) >= 3:
                city_state_zip = f"{city_parts[0]}, {city_parts[1]} {city_parts[2]}"
            elif len(city_parts) == 2:
                city_state_zip = f"{city_parts[0]}, {city_parts[1]}"
            else:
                city_state_zip = city_parts[0]
        customer_text = "\n".join(filter(None, [name_line, address_line, city_state_zip]))
        try:
            cell_a3 = sheet.cell(row=3, column=1)
            if not isinstance(cell_a3, MergedCell):
                cell_a3.value = customer_text
            else:
                merged_range_a3 = None
                merge_range_obj_a3 = None
                for merge_range in list(sheet.merged_cells.ranges):
                    if "A3" in str(merge_range):
                        merged_range_a3 = str(merge_range)
                        merge_range_obj_a3 = merge_range
                        break
                if merged_range_a3:
                    sheet.unmerge_cells(merged_range_a3)
                    sheet.cell(merge_range_obj_a3.min_row, merge_range_obj_a3.min_col).value = customer_text
                    sheet.merge_cells(merged_range_a3)
                else:
                    sheet.cell(row=3, column=1).value = customer_text
        except Exception:
            sheet.cell(row=3, column=1).value = customer_text

    # Serial list and panel count
    items = sorted(pallet.items, key=lambda i: i.slot_index)
    serials = [item.serial for item in items]
    panel_count = len(serials)

    # B1: template/panel type
    if template_type:
        try:
            cell_b1 = sheet.cell(row=1, column=2)
            if not isinstance(cell_b1, MergedCell):
                cell_b1.value = template_type
            else:
                merged_range_b1 = None
                merge_range_obj_b1 = None
                for merge_range in list(sheet.merged_cells.ranges):
                    if "B1" in str(merge_range):
                        merged_range_b1 = str(merge_range)
                        merge_range_obj_b1 = merge_range
                        break
                if merged_range_b1:
                    sheet.unmerge_cells(merged_range_b1)
                    sheet.cell(merge_range_obj_b1.min_row, merge_range_obj_b1.min_col).value = template_type
                    sheet.merge_cells(merged_range_b1)
                else:
                    sheet.cell(row=1, column=2).value = template_type
        except Exception:
            sheet.cell(row=1, column=2).value = template_type

    # B2: panel count
    try:
        cell_b2 = sheet.cell(row=2, column=2)
        if not isinstance(cell_b2, MergedCell):
            cell_b2.value = panel_count
        else:
            merged_range_b2 = None
            merge_range_obj_b2 = None
            for merge_range in list(sheet.merged_cells.ranges):
                if "B2" in str(merge_range):
                    merged_range_b2 = str(merge_range)
                    merge_range_obj_b2 = merge_range
                    break
            if merged_range_b2:
                sheet.unmerge_cells(merged_range_b2)
                sheet.cell(merge_range_obj_b2.min_row, merge_range_obj_b2.min_col).value = panel_count
                sheet.merge_cells(merged_range_b2)
            else:
                sheet.cell(row=2, column=2).value = panel_count
    except Exception:
        sheet.cell(row=2, column=2).value = panel_count

    # D2: weight (panel_count * 40), as in 1.1
    weight = panel_count * 40
    try:
        cell_d2 = sheet.cell(row=2, column=4)
        if not isinstance(cell_d2, MergedCell):
            cell_d2.value = weight
        else:
            merged_range_d2 = None
            merge_range_obj_d2 = None
            for merge_range in list(sheet.merged_cells.ranges):
                if "D2" in str(merge_range):
                    merged_range_d2 = str(merge_range)
                    merge_range_obj_d2 = merge_range
                    break
            if merged_range_d2:
                sheet.unmerge_cells(merged_range_d2)
                sheet.cell(merge_range_obj_d2.min_row, merge_range_obj_d2.min_col).value = weight
                sheet.merge_cells(merged_range_d2)
            else:
                sheet.cell(row=2, column=4).value = weight
    except Exception:
        sheet.cell(row=2, column=4).value = weight

    # G3: date label d-Mmm-yy
    try:
        try:
            date_label = export_dt.strftime("%-d-%b-%y")
        except ValueError:
            date_label = export_dt.strftime("%#d-%b-%y")
        cell_g3 = sheet.cell(row=3, column=7)
        if not isinstance(cell_g3, MergedCell):
            cell_g3.value = date_label
        else:
            merged_range_g3 = None
            merge_range_obj_g3 = None
            for merge_range in list(sheet.merged_cells.ranges):
                if "G3" in str(merge_range):
                    merged_range_g3 = str(merge_range)
                    merge_range_obj_g3 = merge_range
                    break
            if merged_range_g3:
                sheet.unmerge_cells(merged_range_g3)
                sheet.cell(merge_range_obj_g3.min_row, merge_range_obj_g3.min_col).value = date_label
                sheet.merge_cells(merged_range_g3)
            else:
                sheet.cell(row=3, column=7).value = date_label
    except Exception:
        sheet.cell(row=3, column=7).value = export_dt.date().isoformat()

    # B3: panelType + MDYYYY + "-" + pallet_number
    if template_type:
        try:
            try:
                date_mdyyyy = export_dt.strftime("%-m%-d%Y")
            except ValueError:
                date_mdyyyy = export_dt.strftime("%#m%#d%Y")
            b3_value = f"{template_type}{date_mdyyyy}-{pallet.pallet_number}"
            cell_b3 = sheet.cell(row=3, column=2)
            if not isinstance(cell_b3, MergedCell):
                cell_b3.value = b3_value
            else:
                merged_range_b3 = None
                merge_range_obj_b3 = None
                for merge_range in list(sheet.merged_cells.ranges):
                    if "B3" in str(merge_range):
                        merged_range_b3 = str(merge_range)
                        merge_range_obj_b3 = merge_range
                        break
                if merged_range_b3:
                    sheet.unmerge_cells(merged_range_b3)
                    sheet.cell(merge_range_obj_b3.min_row, merge_range_obj_b3.min_col).value = b3_value
                    sheet.merge_cells(merged_range_b3)
                else:
                    sheet.cell(row=3, column=2).value = b3_value
        except Exception:
            sheet.cell(row=3, column=2).value = f"{template_type}-{pallet.pallet_number}"

    # Serial + electricals
    serial_col = "B"
    start_row = 5
    pm_col = _find_column_by_header(sheet, ["Pm", "Pm(W)", "Pm (W)"])
    isc_col = _find_column_by_header(sheet, ["Isc", "Isc(A)", "Isc (A)"])
    voc_col = _find_column_by_header(sheet, ["Voc", "Voc(V)", "Voc (V)", "Voc(V)"])
    ipm_col = _find_column_by_header(sheet, ["Ipm", "Ipm(A)", "Ipm (A)"])
    vpm_col = _find_column_by_header(sheet, ["Vpm", "Vpm(V)", "Vpm (V)", "Vpm(V)"])

    total_serials = len(serials)
    for idx, serial in enumerate(serials):
        row_idx = start_row + idx
        sheet[f"{serial_col}{row_idx}"].value = serial

        data = serial_data_cache.get(serial.upper()) or {}
        pm_val = data.get("Pm")
        use_theoretical = not _validate_pm_range(pm_val, template_type)
        if use_theoretical:
            fallback = _generate_theoretical_electrical_values(template_type)
            pm_val = fallback["Pm"]
            isc_val = fallback["Isc"]
            voc_val = fallback["Voc"]
            ipm_val = fallback["Ipm"]
            vpm_val = fallback["Vpm"]
        else:
            isc_val = data.get("Isc")
            voc_val = data.get("Voc")
            ipm_val = data.get("Ipm")
            vpm_val = data.get("Vpm")

        if pm_col and pm_val is not None:
            sheet[f"{pm_col}{row_idx}"].value = pm_val
        if isc_col and isc_val is not None:
            sheet[f"{isc_col}{row_idx}"].value = isc_val
        if voc_col and voc_val is not None:
            sheet[f"{voc_col}{row_idx}"].value = voc_val
        if ipm_col and ipm_val is not None:
            sheet[f"{ipm_col}{row_idx}"].value = ipm_val
        if vpm_col and vpm_val is not None:
            sheet[f"{vpm_col}{row_idx}"].value = vpm_val

        # Keep loop in place in case we later want progress callbacks.
        _ = total_serials


def write_legacy_excel_export(
    pallet: Pallet,
    template_type: str,
    db: Session,
    *,
    export_dt: datetime | None = None,
    base_dir: Path | None = None,
) -> Path | None:
    """Excel export that mirrors 1.1 PALLET SHEET behavior as closely as possible."""
    try:
        excel_dir = Path("data") / "EXCEL"
        source_workbook = _find_pallet_workbook(excel_dir, pallet.max_panels)
        if source_workbook is None or not source_workbook.exists():
            # Fallback to simple workbook if template not configured.
            root = base_dir or Path("data") / "PALLETS"
            now = datetime.now(timezone.utc)
            try:
                date_label = now.strftime("%-d-%b-%y")
            except ValueError:
                date_label = now.strftime("%d-%b-%y")
            export_dir = root / date_label
            export_dir.mkdir(parents=True, exist_ok=True)
            timestamp = now.strftime("%Y-%m-%d_%H-%M-%S")
            filename = f"Pallet_{pallet.pallet_number:03d}_{timestamp}.xlsx"
            path = export_dir / filename
            wb = Workbook()
            ws = wb.active
            ws.title = "Pallet"
            ws["A1"] = "Pallet Number"
            ws["B1"] = pallet.pallet_number
            wb.save(path)
            return path

        # Date-based export directory (always uses current time so filesystem layout
        # reflects when the export was generated, even if packout_date is backdated).
        root = base_dir or Path("data") / "PALLETS"
        now = datetime.now(timezone.utc)
        try:
            date_label = now.strftime("%-d-%b-%y")
        except ValueError:
            date_label = now.strftime("%#d-%b-%y")
        export_dir = root / date_label
        export_dir.mkdir(parents=True, exist_ok=True)

        temp_path = export_dir / "temp_pallet_export.xlsx"
        # Copy entire workbook
        import shutil as _shutil

        _shutil.copyfile(source_workbook, temp_path)

        wb = load_workbook(temp_path, read_only=False, keep_vba=False, keep_links=False, data_only=False)
        try:
            sheet_name = None
            for name in wb.sheetnames:
                if name.upper().replace(" ", "") == "PALLETSHEET":
                    sheet_name = name
                    break
            if sheet_name is None:
                return None

            # Build serial->electricals cache from SimPanel
            serials = [item.serial for item in pallet.items]
            if serials:
                rows = (
                    db.query(SimPanel)
                    .filter(SimPanel.serial.in_(serials))
                    .order_by(SimPanel.serial.asc(), SimPanel.test_timestamp.desc())
                    .all()
                )
            else:
                rows = []
            serial_data_cache: dict[str, dict] = {}
            for row in rows:
                key = row.serial.upper()
                if key in serial_data_cache:
                    continue
                serial_data_cache[key] = {
                    "Pm": row.pm,
                    "Isc": row.isc,
                    "Voc": row.voc,
                    "Ipm": row.ipm,
                    "Vpm": row.vpm,
                }

            customer = db.query(Customer).filter(Customer.id == pallet.customer_id).first() if pallet.customer_id else None
            sheet = wb[sheet_name]
            # Use explicitly provided export_dt for backdating (packout date) if present;
            # otherwise fall back to "now" so sheets default to today's date.
            sheet_export_dt = export_dt or datetime.now()
            _update_pallet_sheet_from_db(sheet, pallet, template_type, customer, serial_data_cache, sheet_export_dt)

            # Use B3 for filename if possible
            b3_val = sheet["B3"].value
            if b3_val:
                raw = str(b3_val)
                invalid_chars = '<>:"/\\|?*'
                base_name = "".join(c if c not in invalid_chars else "_" for c in raw)
                filename = f"{base_name}.xlsx"
            else:
                timestamp = now.strftime("%Y-%m-%d_%H-%M-%S")
                filename = f"Pallet_{pallet.pallet_number:03d}_{timestamp}.xlsx"

            final_path = export_dir / filename
            counter = 2
            while final_path.exists():
                stem = final_path.stem
                suffix = final_path.suffix
                final_path = export_dir / f"{stem} ({counter}){suffix}"
                counter += 1

            wb.save(final_path)
            return final_path
        finally:
            try:
                wb.close()
            except Exception:
                pass
            try:
                if temp_path.exists():
                    temp_path.unlink()
            except Exception:
                pass
    except Exception:
        return None
